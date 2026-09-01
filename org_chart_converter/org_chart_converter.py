#!/usr/bin/env python3
"""
Convert an org-structure Excel export into a staircase-style org chart
workbook with collapsible leadership grouping. Two source shapes are
auto-detected from the header row:

FORMAT A - one row per manager-chain scope, "Employee" as the leaf column:

    Manager - Level 1 | ... | Employee       | Title  | Department | ...
    All               | ... | All            |        |            |
    N/A               | ... | Aprille Tiedra | Coord. | Support    | ...
    Bob Lyons         | ... | Ben Reich      | VP     | Sales      | ...

  - The column headed exactly "Employee" (case-insensitive) holds an
    individual name for that row. Every column before it forms that
    person's manager chain, read left to right (any number of columns).
    Every column after it is a per-person metadata field, carried through
    to the output as-is, in source order.
  - "All"/"N/A" (any case) and blank cells are filler in the chain columns.
    Rows whose "Employee" cell is filler are rollup/summary rows and are
    skipped - they add no reporting-line information not already present
    elsewhere, and their metadata columns are placeholder junk.
  - A row's parent is the last real name in its chain; an empty chain means
    the person is a root.

FORMAT B - one row per employee, with named ancestor columns going upward:

    employee_id | employee_legal_name | title | department_id | department |
    current_entity | manager_id | manager_legal_name | manager_level_2 | ...
    | top_level_leader

  - "employee_legal_name" is that row's person, renamed "Employee" and put
    first in the output.
  - "manager_legal_name" is their direct manager; "manager_level_2",
    "manager_level_3", etc. are each one generation further up. Blank cells
    end the chain early (that generation has no more managers above them).
  - "top_level_leader" is dropped - it's redundant with whichever chain
    column already holds the top of the org.
  - Any column with "_id" in its name is dropped, and so is "role_state".
  - "department", "current_entity", and "title" become metadata columns at
    the end of the output, in that order. Any other leftover column is
    still carried through (placed before those three) rather than silently
    dropped.

Usage:
    python org_chart_converter.py INPUT.xlsx [-o OUTPUT.xlsx]
"""
import argparse
import re
import sys
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FILLER = {"", "all", "n/a", "na", "none"}
EMPLOYEE_HEADER = "employee"
FORMAT_B_PRIORITY_TRAILING = ["department", "current_entity", "title"]
FORMAT_B_DROPPED_COLUMNS = ["role_state"]


def is_filler(value):
    if value is None:
        return True
    return str(value).strip().lower() in FILLER


def clean(value):
    return str(value).strip()


def _normalize_header(value):
    return re.sub(r"[\s_]+", " ", str(value or "")).strip().lower()


class ParseResult:
    def __init__(self):
        self.parent_of = {}       # name -> parent name (or None for root)
        self.children_of = {}     # name -> ordered list of child names (dedup)
        self.all_names = set()
        self.warnings = []
        self.rows_read = 0
        self.rows_used = 0
        self.attr_headers = []    # ordered metadata column labels (may be empty)
        self.attributes = {}      # name -> {attr_header: value}

    def add_node(self, name):
        self.all_names.add(name)
        self.children_of.setdefault(name, [])

    def set_parent(self, child, parent):
        self.add_node(child)
        if parent is not None:
            self.add_node(parent)

        existing = self.parent_of.get(child, "__unset__")
        if existing == "__unset__":
            self.parent_of[child] = parent
            if parent is not None:
                siblings = self.children_of[parent]
                if child not in siblings:
                    siblings.append(child)
        elif existing != parent:
            self.warnings.append(
                f'"{child}" appears under multiple managers: '
                f'"{existing}" and "{parent}" (kept "{existing}").'
            )

    def set_attributes(self, name, attrs):
        if not attrs:
            return
        existing = self.attributes.get(name)
        if existing is None:
            self.attributes[name] = attrs
        elif existing != attrs:
            self.warnings.append(
                f'"{name}" has conflicting metadata across rows '
                f"(kept the first occurrence)."
            )


def _find_employee_column(header):
    for i, cell in enumerate(header):
        if cell is not None and str(cell).strip().lower() == EMPLOYEE_HEADER:
            return i
    return None


def _find_column(norm_headers, target):
    for i, h in enumerate(norm_headers):
        if h == target:
            return i
    return None


def parse_workbook(path, sheet_name=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Worksheet is empty.")

    header, *data_rows = rows
    if len(header) < 2:
        raise ValueError(
            "Expected at least 2 columns (one or more manager-level "
            "columns plus a final name column)."
        )

    employee_idx = _find_employee_column(header)
    if employee_idx is not None:
        return _parse_format_a(header, data_rows, employee_idx)

    norm_headers = [_normalize_header(h) for h in header]
    if "employee legal name" in norm_headers:
        return _parse_format_b(header, norm_headers, data_rows)

    raise ValueError(
        'Could not find a column headed "Employee" or "employee_legal_name" '
        "in the header row - unrecognized report format."
    )


def _parse_format_a(header, data_rows, employee_idx):
    attr_headers = [
        str(h).strip() for h in header[employee_idx + 1:] if h is not None and str(h).strip()
    ]
    n_attrs = len(attr_headers)

    result = ParseResult()
    result.attr_headers = attr_headers

    for row in data_rows:
        if row is None or all(c is None for c in row):
            continue
        result.rows_read += 1

        level_cells = row[:employee_idx]
        leaf_cell = row[employee_idx]
        attr_cells = row[employee_idx + 1:employee_idx + 1 + n_attrs]

        if is_filler(leaf_cell):
            continue  # rollup/summary row, no new information

        leaf_name = clean(leaf_cell)
        chain = [clean(c) for c in level_cells if not is_filler(c)]

        parent = chain[-1] if chain else None
        result.set_parent(leaf_name, parent)

        if attr_headers:
            attrs = {
                h: (v if v is not None else "")
                for h, v in zip(attr_headers, attr_cells)
            }
            result.set_attributes(leaf_name, attrs)

        # Also register manager-to-manager links within the chain itself,
        # as a fallback for data where a manager never gets its own leaf row.
        for i in range(len(chain) - 1):
            result.set_parent(chain[i + 1], chain[i])

        result.rows_used += 1

    return result


def _parse_format_b(header, norm_headers, data_rows):
    employee_idx = _find_column(norm_headers, "employee legal name")

    # Manager chain, closest-to-furthest from the employee: "manager_legal_name"
    # is generation 1, "manager_level_2" is generation 2, etc.
    chain_idx = []
    direct_manager_idx = _find_column(norm_headers, "manager legal name")
    if direct_manager_idx is not None:
        chain_idx.append((1, direct_manager_idx))
    level_re = re.compile(r"^manager level (\d+)$")
    for i, h in enumerate(norm_headers):
        m = level_re.match(h)
        if m:
            chain_idx.append((int(m.group(1)), i))
    chain_idx.sort(key=lambda pair: pair[0])
    chain_indices = [i for _, i in chain_idx]

    top_leader_idx = _find_column(norm_headers, "top level leader")
    id_indices = {i for i, h in enumerate(header) if h is not None and "_id" in str(h).lower()}
    dropped_indices = {
        _find_column(norm_headers, _normalize_header(name)) for name in FORMAT_B_DROPPED_COLUMNS
    }

    excluded = {employee_idx, top_leader_idx} | id_indices | dropped_indices | set(chain_indices)
    excluded.discard(None)

    priority_indices = []
    for target in FORMAT_B_PRIORITY_TRAILING:
        idx = _find_column(norm_headers, _normalize_header(target))
        if idx is not None and idx not in excluded:
            priority_indices.append(idx)
    other_indices = [
        i for i in range(len(header))
        if i not in excluded and i not in priority_indices
    ]
    attr_indices = other_indices + priority_indices
    attr_headers = [str(header[i]).strip() for i in attr_indices]

    result = ParseResult()
    result.attr_headers = attr_headers

    for row in data_rows:
        if row is None or all(c is None for c in row):
            continue
        result.rows_read += 1

        employee_cell = row[employee_idx]
        if is_filler(employee_cell):
            continue

        employee_name = clean(employee_cell)
        chain = [clean(row[i]) for i in chain_indices if not is_filler(row[i])]

        parent = chain[0] if chain else None
        result.set_parent(employee_name, parent)

        # Manager-to-manager links, as a fallback for a manager who never
        # gets their own "employee_legal_name" row in this export.
        for i in range(len(chain) - 1):
            result.set_parent(chain[i], chain[i + 1])

        if attr_headers:
            attrs = {
                h: (row[i] if row[i] is not None else "")
                for h, i in zip(attr_headers, attr_indices)
            }
            result.set_attributes(employee_name, attrs)

        result.rows_used += 1

    return result


def build_forest(result):
    roots = [
        name for name in result.all_names
        if result.parent_of.get(name) is None
    ]
    roots.sort(key=str.lower)
    for parent in result.children_of:
        result.children_of[parent].sort(key=str.lower)
    return roots


def max_depth(result, roots):
    best = 0

    def walk(name, depth):
        nonlocal best
        best = max(best, depth)
        for child in result.children_of.get(name, []):
            walk(child, depth + 1)

    for r in roots:
        walk(r, 0)
    return best


HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=12)
MANAGER_FONT = Font(bold=True)
LEAF_FONT = Font(bold=False)
ROOT_FONT = Font(bold=True, size=13, color="1F4E78")


def write_org_chart(result, roots, out_path, source_name=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "Org Chart"

    depth = max_depth(result, roots)
    staircase_cols = depth + 1
    attr_headers = result.attr_headers
    # Column 1 is a fixed "Employee" column (always shows this row's person,
    # regardless of which staircase column their depth puts them in).
    # Columns 2..staircase_cols+1 are the staircase itself.
    name_col = 1
    staircase_start = 2
    attrs_start = staircase_start + staircase_cols
    total_cols = staircase_cols + 1 + len(attr_headers)

    ws.cell(row=1, column=name_col, value="Employee")
    ws["A1"].font = HEADER_FONT
    ws["A1"].fill = HEADER_FILL
    ws.cell(row=1, column=staircase_start,
             value=f"Org chart{(' - ' + source_name) if source_name else ''}")
    for col in range(staircase_start, attrs_start):
        ws.cell(row=1, column=col).fill = HEADER_FILL
    ws.cell(row=1, column=staircase_start).font = HEADER_FONT

    for i, label in enumerate(attr_headers):
        cell = ws.cell(row=1, column=attrs_start + i, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    ws.row_dimensions[1].height = 22

    current_row = 2

    def write_node(name, col, depth_level):
        nonlocal current_row
        r = current_row
        ws.cell(row=r, column=name_col, value=name).font = Font(bold=(depth_level == 0))
        cell = ws.cell(row=r, column=col, value=name)
        children = result.children_of.get(name, [])
        if depth_level == 0:
            cell.font = ROOT_FONT
        elif children:
            cell.font = MANAGER_FONT
        else:
            cell.font = LEAF_FONT

        if depth_level > 0:
            ws.row_dimensions[r].outline_level = min(depth_level, 7)

        attrs = result.attributes.get(name)
        if attrs:
            for i, header in enumerate(attr_headers):
                ws.cell(row=r, column=attrs_start + i, value=attrs.get(header, ""))

        current_row += 1
        for child in children:
            write_node(child, col + 1, depth_level + 1)

    for root in roots:
        write_node(root, staircase_start, 0)

    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.column_dimensions[get_column_letter(name_col)].width = 26
    for col in range(staircase_start, total_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 26
    ws.freeze_panes = "B2"

    # Summary sheet
    summary = wb.create_sheet("Summary")
    summary["A1"] = "Metric"
    summary["B1"] = "Value"
    summary["A1"].font = summary["B1"].font = Font(bold=True)
    stats = [
        ("Source file", source_name),
        ("Rows read", result.rows_read),
        ("Rows used", result.rows_used),
        ("Rows skipped (rollup/summary)", result.rows_read - result.rows_used),
        ("Total people", len(result.all_names)),
        ("Top-level roots", len(roots)),
        ("Max depth", depth + 1),
        ("Metadata columns detected", ", ".join(attr_headers) if attr_headers else "(none)"),
        ("Data warnings", len(result.warnings)),
    ]
    for i, (k, v) in enumerate(stats, start=2):
        summary.cell(row=i, column=1, value=k)
        summary.cell(row=i, column=2, value=v)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 40

    if result.warnings:
        summary.cell(row=len(stats) + 3, column=1, value="Warnings").font = Font(bold=True)
        for i, w in enumerate(result.warnings, start=len(stats) + 4):
            summary.cell(row=i, column=1, value=w)

    wb.save(out_path)


def convert(input_path, output_path=None, sheet_name=None):
    result = parse_workbook(input_path, sheet_name=sheet_name)
    roots = build_forest(result)
    if not roots:
        raise ValueError("No root (top-of-org) people found - check the input format.")
    if output_path is None:
        output_path = input_path.rsplit(".", 1)[0] + "_wireframe.xlsx"
    write_org_chart(result, roots, output_path, source_name=input_path.split("/")[-1])
    return result, roots, output_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="Path to the source .xlsx report")
    parser.add_argument("-o", "--output", help="Path to write the wireframe .xlsx to")
    parser.add_argument("--sheet", help="Worksheet name to read (defaults to the active sheet)")
    args = parser.parse_args()

    try:
        result, roots, output_path = convert(args.input, args.output, args.sheet)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

    print(f"Read {result.rows_read} rows, used {result.rows_used}.")
    print(f"Found {len(result.all_names)} people, {len(roots)} top-level root(s), "
          f"max depth {max_depth(result, roots) + 1}.")
    if result.attr_headers:
        print(f"Metadata columns carried through: {', '.join(result.attr_headers)}")
    if result.warnings:
        print(f"{len(result.warnings)} data warning(s) - see the Summary sheet.")
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    main()
